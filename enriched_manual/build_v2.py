from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
import json, os, sys

# Screenshot mapping: H1 substring -> list of PNG filename queries
SCREENSHOT_MAP = {
    'SITHEHR': ['01-login-page.png'],
    'Table of Contents': [],
    '1. System Architecture': ['01-login-page.png', '03-dashboard.png'],
    '2. Getting Started': ['01-login-page.png', '02-credentials-filled.png', '03-dashboard.png'],
    '3. Master Data Management': [
        'masters-department-01-list.png',
        'masters-users-01-list.png',
        'masters-province-01-list.png',
        'masters-barangay-01-list.png',
        'masters-city-01-list.png',
        'masters-country-01-list.png',
        'masters-nationality-01-list.png',
        'masters-occupation-01-list.png',
        'masters-permissions-01-list.png',
        'masters-hospital-01-list.png',
        'masters-tenant-01-list.png',
    ],
    '4. Patient Registration': ['patients-01-list.png'],
    '5. Public Appointment Booking': ['public-booking-01-list.png'],
    '6. Outpatient Department': [
        'opd-consultations-01-list.png',
        'opd-clinics-01-list.png',
        'opd-services-01-list.png',
        'opd-visit-types-01-list.png',
        'opd-complaints-01-list.png',
        'opd-consultation-types-01-list.png',
        'opd-referral-sources-01-list.png',
        'konsulta-01-list.png',
    ],
    '7. Emergency Room': [
        'er-triage-01-list.png',
        'er-beds-01-list.png',
        'er-zones-01-list.png',
        'er-area-01-list.png',
        'er-acuity-levels-01-list.png',
        'er-mobility-statuses-01-list.png',
    ],
    '8. Inpatient Admission': [
        'admission-01-list.png',
        'wards-01-list.png',
        'ward-rooms-01-list.png',
        'ward-room-types-01-list.png',
        'icu-wards-01-list.png',
        'icu-beds-01-list.png',
        'or-rooms-01-list.png',
        'or-schedule-01-list.png',
        'delivery-rooms-01-list.png',
        'delivery-types-01-list.png',
        'delivery-schedule-01-list.png',
        'delivery-complications-01-list.png',
    ],
    '9. Laboratory Services': [
        'lab-tests-01-list.png',
        'lab-results-01-list.png',
        'lab-sections-01-list.png',
        'lab-panels-01-list.png',
        'lab-specimen-types-01-list.png',
        'lab-units-01-list.png',
        'lab-instruments-01-list.png',
        'lab-reference-ranges-01-list.png',
        'lab-dispensing-01-list.png',
    ],
    '10. Radiology Services': [
        'radiology-results-01-list.png',
        'radiology-modalities-01-list.png',
        'radiology-body-parts-01-list.png',
        'radiology-exam-types-01-list.png',
        'radiology-dispensing-01-list.png',
    ],
    '11. Health Sciences': [
        'pharmacy-drugs-01-list.png',
        'pharmacy-inventory-01-list.png',
        'pharmacy-dispensing-01-list.png',
        'pharmacy-purchase-orders-01-list.png',
        'pharmacy-drug-batches-01-list.png',
        'pharmacy-stock-adjustments-01-list.png',
        'pharmacy-patient-record-01-list.png',
        'pharmacy-reports-01-list.png',
        'supply-items-01-list.png',
        'supply-categories-01-list.png',
        'supply-units-01-list.png',
        'supply-stock-locations-01-list.png',
        'supply-suppliers-01-list.png',
        'supply-inventory-01-list.png',
        'supply-dispensing-01-list.png',
        'supply-stock-adjustments-01-list.png',
        'procedure-procedures-01-list.png',
        'procedure-categories-01-list.png',
        'procedure-dispensing-01-list.png',
    ],
    '12. Billing and Payments': [
        'billing-01-list.png',
        'billing-fund-sources-01-list.png',
        'billing-insurance-providers-01-list.png',
        'dashboard-01-list.png',
    ],
    '13. PhilHealth Module': [
        'philhealth-members-01-list.png',
        'philhealth-claims-01-list.png',
        'philhealth-icd10-01-list.png',
        'philhealth-case-rates-01-list.png',
        'philhealth-batches-01-list.png',
        'philhealth-rvs-01-list.png',
        'philhealth-hc-professionals-01-list.png',
        'philhealth-reports-01-list.png',
        'yakap-dashboard-01-list.png',
        'yakap-empanelments-01-list.png',
        'yakap-screenings-01-list.png',
        'yakap-fpe-01-list.png',
        'yakap-capitation-01-list.png',
    ],
    '14. Medical Records': [
        'medical-records-01-list.png',
    ],
    '15. Doctors Module': [
        'doctors-dashboard-01-list.png',
        'billing-01-list.png',
    ],
    '16. OB, Delivery': [
        'ob-visit-types-01-list.png',
        'ob-consultations-01-list.png',
        'ob-fetal-presentations-01-list.png',
        'ob-contraceptive-methods-01-list.png',
    ],
    '17. Nursing': [
        'nursing-01-list.png',
    ],
    '18. Procedures, Delivery': [
        'procedure-procedures-01-list.png',
        'procedure-categories-01-list.png',
        'procedure-dispensing-01-list.png',
        'delivery-types-01-list.png',
        'delivery-rooms-01-list.png',
    ],
    '19. Frequently Asked Questions': [],
    '20. Appendix': [],
}

PNG_SOURCES = [
    'E:/sithemr/module-screenshots',
    'E:/sithemr/opd-screenshots',
]

def get_png_registry():
    pngs = {}
    for src in PNG_SOURCES:
        if os.path.isdir(src):
            for f in os.listdir(src):
                if f.lower().endswith('.png'):
                    pngs[f.lower()] = f
    return pngs

PNG_REGISTRY = get_png_registry()

def find_png(query):
    q = query.lower()
    if q in PNG_REGISTRY:
        return PNG_REGISTRY[q]
    for k, v in PNG_REGISTRY.items():
        if q in k or k in q:
            return v
    return None

def add_screenshot(doc, query, caption='', width_mm=130):
    png_name = find_png(query)
    if not png_name:
        return False
    for src in PNG_SOURCES:
        fp = os.path.join(src, png_name)
        if os.path.isfile(fp):
            break
    else:
        return False
    try:
        with Image.open(fp) as img:
            aspect = img.size[1] / img.size[0]
        p = doc.add_paragraph()
        p.alignment = 1
        p.add_run().add_picture(fp, width=Mm(width_mm))
        if caption:
            cp = doc.add_paragraph()
            cp.alignment = 1
            cr = cp.add_run(caption)
            cr.font.size = Pt(8)
            cr.font.italic = True
            cr.font.color.rgb = RGBColor(100, 100, 100)
        doc.add_paragraph()
        return True
    except:
        return False

def screenshot_for_section(doc, h1_text):
    inserted = 0
    for key, imgs in SCREENSHOT_MAP.items():
        if key in h1_text:
            for img in imgs:
                if add_screenshot(doc, img, 'Screenshot: ' + img, 130):
                    inserted += 1
    return inserted

def build_manual(json_files, output_path):
    doc = Document()
    for section in doc.sections:
        section.page_height = Mm(279.4)
        section.page_width = Mm(215.9)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(18)
        section.right_margin = Mm(18)
    
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15
    
    for lvl in range(1, 4):
        hs = doc.styles['Heading %d' % lvl]
        hs.font.color.rgb = RGBColor(0, 51, 102)
        hs.font.name = 'Calibri'
        if lvl == 1:
            hs.font.size = Pt(18)
            hs.paragraph_format.space_before = Pt(24)
        elif lvl == 2:
            hs.font.size = Pt(14)
        else:
            hs.font.size = Pt(12)
    
    bp = 0
    se = 0
    for jf in json_files:
        if not os.path.isfile(jf):
            print('SKIP missing: ' + jf)
            continue
        with open(jf, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        is_list = isinstance(data, list)
        blocks = data if is_list else data.get('blocks', [])
        header_txt = data.get('header', '') if isinstance(data, dict) else ''
        
        section = doc.sections[-1]
        if header_txt:
            hdr = section.header
            hdr.is_linked_to_previous = False
            hp = hdr.paragraphs[0]
            hp.alignment = 1
            hr = hp.add_run(header_txt)
            hr.font.size = Pt(8)
            hr.font.italic = True
            hr.font.color.rgb = RGBColor(128, 128, 128)
        
        for block in blocks:
            bt = block.get('type', '')
            bp += 1
            if bt == 'heading':
                lvl = block.get('level', 1)
                txt = block.get('text', '')
                if lvl == 1 and txt:
                    ins = screenshot_for_section(doc, txt)
                    if ins:
                        se += ins
                doc.add_heading(txt, level=lvl)
            elif bt == 'paragraph':
                txt = block.get('text', '')
                if txt:
                    doc.add_paragraph(txt)
            elif bt == 'runs':
                p = doc.add_paragraph()
                for rd in block.get('runs', []):
                    r = p.add_run(rd.get('text', ''))
                    if rd.get('bold'): r.bold = True
                    if rd.get('italic'): r.italic = True
                    if rd.get('size'): r.font.size = Pt(rd['size'])
            elif bt == 'bullet_list':
                for item in block.get('items', []):
                    doc.add_paragraph(item, style='List Bullet')
            elif bt == 'numbered_list':
                for item in block.get('items', []):
                    doc.add_paragraph(item, style='List Number')
            elif bt == 'table':
                hdrs = block.get('header', [])
                rows = block.get('rows', [])
                if not hdrs and rows:
                    hdrs = rows[0]
                    rows = rows[1:]
                ncols = len(hdrs) if hdrs else (len(rows[0]) if rows else 1)
                tbl = doc.add_table(rows=1+len(rows), cols=ncols)
                tbl.style = 'Light Grid Accent 1'
                if hdrs:
                    for i, h in enumerate(hdrs):
                        c = tbl.rows[0].cells[i]
                        c.text = ''
                        rn = c.paragraphs[0].add_run(h)
                        rn.bold = True
                        rn.font.size = Pt(10)
                        rn.font.color.rgb = RGBColor(255, 255, 255)
                        shd = OxmlElement('w:shd')
                        shd.set(qn('w:fill'), '003366')
                        c._tc.get_or_add_tcPr().append(shd)
                for ri, row in enumerate(rows):
                    for ci, val in enumerate(row):
                        if ci < ncols:
                            c = tbl.rows[ri+1].cells[ci]
                            c.text = str(val) if val is not None else ''
                            for pp in c.paragraphs:
                                for rn in pp.runs:
                                    rn.font.size = Pt(10)
            elif bt == 'page_break':
                doc.add_page_break()
            elif bt == 'spacer':
                doc.add_paragraph()
            elif bt == 'screenshot':
                png = block.get('png', '')
                cap = block.get('caption', '')
                w = block.get('width_mm', 130)
                if add_screenshot(doc, png, cap, w):
                    se += 1
    
    for section in doc.sections:
        ftr = section.footer
        ftr.is_linked_to_previous = False
        fp = ftr.paragraphs[0]
        fp.alignment = 1
        r1 = fp.add_run('Page ')
        r1.font.size = Pt(8)
        r1.font.color.rgb = RGBColor(128, 128, 128)
        fld1 = OxmlElement('w:fldChar')
        fld1.set(qn('w:fldCharType'), 'begin')
        r1._r.append(fld1)
        ins1 = OxmlElement('w:instrText')
        ins1.set(qn('xml:space'), 'preserve')
        ins1.text = ' PAGE '
        r1._r.append(ins1)
        fld2 = OxmlElement('w:fldChar')
        fld2.set(qn('w:fldCharType'), 'end')
        r1._r.append(fld2)
        r2 = fp.add_run(' of ')
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGBColor(128, 128, 128)
        fld3 = OxmlElement('w:fldChar')
        fld3.set(qn('w:fldCharType'), 'begin')
        r2._r.append(fld3)
        ins2 = OxmlElement('w:instrText')
        ins2.set(qn('xml:space'), 'preserve')
        ins2.text = ' NUMPAGES '
        r2._r.append(ins2)
        fld4 = OxmlElement('w:fldChar')
        fld4.set(qn('w:fldCharType'), 'end')
        r2._r.append(fld4)
    
    doc.save(output_path)
    return {'ok': True, 'output': os.path.basename(output_path), 'blocks': bp, 'screenshots': se, 'size': os.path.getsize(output_path)}

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print('Usage: build_v2.py output.docx part1.json [part2.json ...]')
        sys.exit(1)
    output = sys.argv[1]
    json_files = sys.argv[2:]
    print('PNG registry: %d images' % len(PNG_REGISTRY))
    print('Building from %d JSON parts...' % len(json_files))
    result = build_manual(json_files, output)
    print('Done:', result)
